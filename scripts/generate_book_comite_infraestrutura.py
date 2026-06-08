from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER_TYPE
from pptx.util import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, ListFlowable, ListItem, PageBreak

TEMPLATE_PATH = 'docs/templetes/Plano_Diretor_Infra_Update_Jun26_vfinal.pptx'
OUTPUT_PPTX = 'reports/book_comite_infraestrutura.pptx'
OUTPUT_PDF = 'reports/book_comite_infraestrutura.pdf'

slides = [
    {'layout': 0, 'title': 'ALE COMBUSTÍVEIS · DSTL', 'subtitle': 'Book de Direcionamento – Plano Diretor de Infraestrutura'},
    {'layout': 2, 'title': 'Agenda do Book', 'bullets': ['1. Síntese executiva e tese de infraestrutura', '2. Diagnóstico de malha e frentes estruturais', '3. Bases próprias, terceiros, monetização, expansão seletiva e regulação', '4. Pipeline, dependências críticas e próxima pauta do comitê', '5. Visão de decisão e recomendações executivas']},
    {'layout': 1, 'title': 'Síntese Executiva', 'bullets': ['Infraestrutura é condição necessária para crescimento sustentável, defesa de margem e redução de dependências estratégicas.', 'O Comitê deve priorizar frentes decisórias, não projetos isolados.', 'A agenda deve proteger bases próprias com gargalos, limitar terceiros e manter vigilância regulatória.']},
    {'layout': 2, 'title': 'Como usar este book', 'bullets': ['O material foi construído por blocos de leitura e frentes de decisão.', 'Cada bloco conecta diagnóstico à decisão requerida do comitê.', 'A base de conteúdo é documental e inclui direcionamentos manuais da responsável pelo projeto.']},
    {'layout': 2, 'title': 'Estrutura do material', 'bullets': ['Bloco 1: Plano de Expansão comercial', 'Bloco 2: Diagnóstico de malha atual', 'Bloco 3: Frentes de decisão estruturais', 'Bloco 4: Pipeline e monitoramento', 'Bloco 5: Visão 2030 e roadmap']},
    {'layout': 2, 'title': 'Bloco 1: Plano de Expansão', 'bullets': ['Define onde a ALE cresce, defende share ou desinveste.', 'A régua comercial orienta cada decisão de infraestrutura.', 'O plano de expansão é a primeira régua de validação.']},
    {'layout': 1, 'title': 'Régua comercial e decisão', 'bullets': ['Infraestrutura sem ambição comercial é custo; ambição comercial sem infraestrutura é promessa.', 'Regiões prioritárias devem ser protegidas com ações estruturais.', 'Áreas de defesa exigem manutenção de participação e resiliência.']},
    {'layout': 2, 'title': 'Bloco 2: Diagnóstico da malha atual', 'bullets': ['Identifica gargalos, capacidades e dependências da rede.', 'Mostra onde a ALE precisa agir para manter oferta e margem.', 'Serve de base para priorizar frentes em vez de projetos isolados.']},
    {'layout': 1, 'title': 'Diagnóstico de malha atual', 'bullets': ['Bases próprias com restrição de capacidade exigem ação urgente.', 'Terceiros e contratos críticos reduzem flexibilidade comercial.', 'A área de influência e as variáveis regulatórias moldam a decisão.']},
    {'layout': 2, 'title': 'Bloco 3: Frentes de decisão', 'bullets': ['Bases próprias com gargalos estruturais', 'Dependência de terceiros', 'Monetização e parcerias', 'Expansão seletiva', 'Riscos regulatórios e transição de produto']},
    {'layout': 1, 'title': 'Bases próprias com gargalos', 'bullets': ['São José do Rio Preto e Betim são pontos de capacidade crítica.', 'Brasília deve ser monitorada como área de influência.', 'Ação agora evita que ativos permaneçam em holding ou sem desempenho operacional.']},
    {'layout': 1, 'title': 'São José do Rio Preto', 'bullets': ['Déficit crítico de capacidade e infraestrutura.', 'O projeto em holding precisa sair da fase de avaliação.', 'Decisão requerida: aprovar escopo técnico e financeiro da intervenção.']},
    {'layout': 1, 'title': 'Betim', 'bullets': ['Plano de revamp hidráulico e automação está em desenvolvimento.', 'Objetivo: aumentar capacidade de 90 mil m³/mês para 140 mil m³/mês.', 'Decisão requerida: autorizar cronograma e recursos para execução.']},
    {'layout': 1, 'title': 'Brasília e área de influência', 'bullets': ['Não há opções relevantes de armazenagem na região.', 'A região deve ser avaliada por rentabilidade e influência.', 'Manter Brasília no radar do comitê como área estratégica.']},
    {'layout': 1, 'title': 'Açailândia / LEM', 'bullets': ['Formalização ANP é condição para consolidar o corredor MATOPIBA.', 'A frente é parte da defesa de infraestrutura crítica.', 'Decisão requerida: aprovar a agenda de formalização e acompanhamento.']},
    {'layout': 2, 'title': 'Dependência crítica de terceiros', 'bullets': ['Terceiros geram custos de disponibilidade e risco de continuidade.', 'Vitória, Duque de Caxias, Goiânia e Guamaré são frentes chave.', 'A decisão deve reduzir dependências e ampliar controle estratégico.']},
    {'layout': 1, 'title': 'Vitória: Atlântica vs SPE', 'bullets': ['Definir o modelo correto para reduzir dependência de Oiltanking.', 'Escolha de alternativa impacta margem e disponibilidade.', 'Decisão requerida: aprovar a direção estratégica e parâmetros de transição.']},
    {'layout': 1, 'title': 'Duque de Caxias', 'bullets': ['Dependência de S10 via cessão de espaço com Raízen aumenta custo.', 'A base não possui duto interligado ao S10; o atual segue dedicado a Marítimo/S500.', 'Decisão requerida: avaliar o modelo atual e potencial interligação direta.']},
    {'layout': 1, 'title': 'Goiânia / Nexta', 'bullets': ['A parceria com Nexta deve ser complementada por solução estrutural.', 'Proposta de dois tanques de 1.500 m³ para reduzir vulnerabilidade.', 'Decisão requerida: autorizar termos de arrendamento e estrutura de tanques.']},
    {'layout': 1, 'title': 'Guamaré', 'bullets': ['Perda de área de influência: -30% de movimentação ALE entre 2024 e 2026.', 'Saída da Raízen e entrada do Pecém ampliam o risco estrutural.', 'Em 2026, 47,5% do estado foi abastecido por venda direta de outros estados.']},
    {'layout': 2, 'title': 'Monetização e parcerias', 'bullets': ['Ativos ociosos devem ser monetizados antes de se tornarem custo.', 'Guarulhos+ é um ativo estratégico de bundling.', 'Santa Maria deve ser mantida como alternativa comercial sem CAPEX.']},
    {'layout': 1, 'title': 'Guarulhos+ Bundling', 'bullets': ['Modelo de bundling deve capturar valor de parcerias estratégicas.', 'Players de interesse: INPASA, FS, DTC, Ipiranga, Nimofast, Midas e outros.', 'Decisão requerida: aprovar o modelo e os parceiros prioritários.']},
    {'layout': 1, 'title': 'Santa Maria', 'bullets': ['Preservar Santa Maria como alternativa comercial sem CAPEX.', 'Não transformar o ativo em projeto de capital direto.', 'Decisão requerida: manter a posição comercial aberta.']},
    {'layout': 1, 'title': 'Novas ideias e oportunidades', 'bullets': ['Identificar oportunidades geradoras de receita e margem.', 'Focar em cases que tragam capital integrado e parcerias comerciais.', 'Manter o pipeline de inovação alinhado ao Plano de Expansão.']},
    {'layout': 2, 'title': 'Expansão seletiva', 'bullets': ['A expansão deve ser seletiva e orientada por valor.', 'Pecém e Itajaí são frentes prioritárias de greenfield.', 'A decisão deve separar crescimento por demanda de iniciativas especulativas.']},
    {'layout': 1, 'title': 'Pecém Greenfield', 'bullets': ['A sondagem Pecém segue como tema de alinhamento com o mercado.', 'Validar a continuidade da avaliação greenfield.', 'A oportunidade deve ser acompanhada sem expectativa de ação imediata.']},
    {'layout': 1, 'title': 'Itajaí / Grupo Ávila', 'bullets': ['A solução atual é um terminal greenfield.', 'O foco deve ser o alinhamento com o Grupo Ávila.', 'Decisão requerida: confirmar a posição de Itajaí como desenvolvimento seletivo.']},
    {'layout': 1, 'title': 'Pipeline de novos negócios', 'bullets': ['Monitorar projetos em análise e em execução.', 'Manter o pipeline de locais não rentáveis com revisão periódica.', 'Conectar oportunidades a decisões rápidas do comitê.']},
    {'layout': 2, 'title': 'Defesa de área de influência', 'bullets': ['Defender áreas estratégicas é tão importante quanto expandir capacidade.', 'Guamaré e Brasília exigem vigilância de influência e rentabilidade.', 'A perda de área pode agravar o impacto de novos concorrentes.']},
    {'layout': 2, 'title': 'Riscos regulatórios', 'bullets': ['Temas regulatórios podem bloquear operação e manter ativos em holding.', 'Cuiabá/Várzea Grande e transição S500 → S10 exigem vigilância.', 'Decisão requerida: manter os temas como monitoramento estratégico.']},
    {'layout': 1, 'title': 'Cuiabá / Várzea Grande', 'bullets': ['A legislação do MT exige comprovação de tancagem própria para operação.', 'O pool deve ser mantido no radar enquanto a consulta ao estado evolui.', 'Decisão requerida: acompanhar a posição regulatória e o modelo equivalente ao país.']},
    {'layout': 1, 'title': 'Transição S500 → S10', 'bullets': ['A transição afeta dutos, bases e produto marítimo.', 'Não há ação estrutural até que haja clareza regulatória e de mercado.', 'Decisão requerida: manter o tema no radar e aguardar confirmação documental.']},
    {'layout': 1, 'title': 'Dependências críticas', 'bullets': ['Formalização ANP para Açailândia / LEM.', 'Termos de Duque de Caxias / Raízen e interligação S10.', 'Modelo de Vitória e termos em Goiânia / Nexta.', 'Acordos de bundling e parceiros em Guarulhos+.']},
    {'layout': 1, 'title': 'Próximos passos objetivos', 'bullets': ['Submeter ação estrutural para São José do Rio Preto e revamp de Betim.', 'Apresentar opções de Vitória e modelo de Duque de Caxias ao comitê.', 'Levar ao comitê a continuidade da sondagem Pecém e a posição de Itajaí.', 'Manter Cuiabá/Várzea Grande e S500 → S10 como itens de acompanhamento.']},
    {'layout': 1, 'title': 'Recomendações para a pauta do comitê', 'bullets': ['Decidir sobre bases críticas, dependência de terceiros e monetização.', 'Priorizar frentes estruturais em vez de projetos isolados.', 'Garantir recursos e governança para execução imediata.']},
    {'layout': 2, 'title': 'Visão 2030 / Roadmap de decisões', 'bullets': ['Consolidar a malha ALE em 2030 com decisões escalonadas.', 'Usar o comitê para manter ritmo de execução e monitoramento.', 'Transformar o plano diretor em trilha clara de decisões.']},
    {'layout': 1, 'title': 'Resumo de tese e prioridades', 'bullets': ['Infraestrutura é condição de competitividade e defesa de margem.', 'Focar em bases próprias, terceiros, monetização, expansão seletiva e regulação.', 'A decisão imediata deve ser robusta, prática e documentada.']},
    {'layout': 5, 'title': 'Confidencial · Uso Interno ALE', 'bullets': []},
]


def add_slide(prs, slide_def):
    slide = prs.slides.add_slide(prs.slide_layouts[slide_def['layout']])
    if slide_def.get('title'):
        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.type in (PP_PLACEHOLDER_TYPE.TITLE, PP_PLACEHOLDER_TYPE.CENTER_TITLE):
                    shape.text = slide_def['title']
                    break
            except Exception:
                continue
    if slide_def.get('subtitle'):
        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.type == PP_PLACEHOLDER_TYPE.SUBTITLE:
                    shape.text = slide_def['subtitle']
                    break
            except Exception:
                continue
    if slide_def.get('bullets'):
        body_shape = None
        for shape in slide.placeholders:
            try:
                if shape.placeholder_format.type in (PP_PLACEHOLDER_TYPE.BODY, PP_PLACEHOLDER_TYPE.OBJECT):
                    body_shape = shape
                    break
            except Exception:
                continue
        if body_shape is None:
            for shape in slide.shapes:
                if shape.has_text_frame and shape is not slide.shapes.title:
                    body_shape = shape
                    break
        if body_shape is not None and body_shape.has_text_frame:
            tf = body_shape.text_frame
            tf.clear()
            for idx, bullet in enumerate(slide_def['bullets']):
                p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
                p.text = bullet
                p.level = 0
                p.font.size = Pt(16)
                p.font.name = 'Arial'
                p.font.bold = False
    return slide


def build_presentation():
    prs = Presentation(TEMPLATE_PATH)
    slide_ids = list(prs.slides._sldIdLst)
    for slide_id in slide_ids:
        prs.slides._sldIdLst.remove(slide_id)
    for slide_def in slides:
        add_slide(prs, slide_def)
    prs.save(OUTPUT_PPTX)
    return OUTPUT_PPTX


def build_pdf():
    doc = SimpleDocTemplate(OUTPUT_PDF, pagesize=A4,
                            rightMargin=40, leftMargin=40,
                            topMargin=40, bottomMargin=40)
    stylesheet = getSampleStyleSheet()
    title_style = ParagraphStyle('Title', parent=stylesheet['Heading1'], fontSize=18, leading=22)
    bullet_style = ParagraphStyle('Bullet', parent=stylesheet['BodyText'], fontSize=11, leading=14, leftIndent=12)
    elems = []
    for idx, slide in enumerate(slides):
        elems.append(Paragraph(slide['title'], title_style))
        elems.append(Spacer(1, 6))
        if slide.get('bullets'):
            list_items = [ListItem(Paragraph(item, bullet_style), bulletColor='black') for item in slide['bullets']]
            elems.append(ListFlowable(list_items, bulletType='bullet', start='disc', leftIndent=12))
        elems.append(Spacer(1, 16))
        if idx < len(slides) - 1:
            elems.append(PageBreak())
    doc.build(elems)
    return OUTPUT_PDF


if __name__ == '__main__':
    build_presentation()
    build_pdf()
    print('Generated:', OUTPUT_PPTX, OUTPUT_PDF)
